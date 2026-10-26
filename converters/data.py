"""Табличні дані: CSV, TSV, JSON, XLSX, XML -> CSV, TSV, JSON, XLSX, XML, HTML, Markdown."""
from __future__ import annotations

import csv
import html as html_lib
import io
import json
import re
import xml.etree.ElementTree as ET
from pathlib import Path

from core import file_ext, normalize_ext
from converters.documents import read_text_file

Table = tuple[list[str], list[list]]   # (заголовки, рядки)


# --- Читання ------------------------------------------------------------------

def _from_records(records: list[dict]) -> Table:
    headers: list[str] = []
    for rec in records:
        for k in rec:
            if k not in headers:
                headers.append(k)
    rows = [[rec.get(h, "") for h in headers] for rec in records]
    return headers, rows


def _flatten(value, prefix="", out=None) -> dict:
    """{"a": {"b": 1}} -> {"a.b": 1}; списки зберігаються як JSON-рядок."""
    out = {} if out is None else out
    if isinstance(value, dict):
        for k, v in value.items():
            _flatten(v, f"{prefix}.{k}" if prefix else str(k), out)
    elif isinstance(value, list):
        out[prefix or "value"] = json.dumps(value, ensure_ascii=False)
    else:
        out[prefix or "value"] = value
    return out


def read_json(path: Path) -> Table:
    data = json.loads(read_text_file(path))
    if isinstance(data, dict):
        # {"items": [...]} — беремо перший список записів
        lists = [v for v in data.values() if isinstance(v, list) and v and isinstance(v[0], dict)]
        if lists:
            data = lists[0]
        elif all(isinstance(v, list) for v in data.values()):
            # стовпцевий формат {"col": [..], ...}
            keys = list(data)
            n = max((len(v) for v in data.values()), default=0)
            return keys, [[data[k][i] if i < len(data[k]) else "" for k in keys] for i in range(n)]
        else:
            data = [data]
    if isinstance(data, list):
        if data and all(isinstance(r, list) for r in data):
            return [str(h) for h in data[0]], [list(r) for r in data[1:]]
        return _from_records([_flatten(r) for r in data])
    return ["value"], [[data]]


def read_csv(path: Path, delimiter: str | None = None) -> Table:
    text = read_text_file(path)
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return [], []
    return rows[0], rows[1:]


def read_xlsx(path: Path) -> Table:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = [["" if v is None else v for v in row] for row in ws.iter_rows(values_only=True)]
    wb.close()
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return [], []
    return [str(h) for h in rows[0]], rows[1:]


def read_xml(path: Path) -> Table:
    root = ET.parse(path).getroot()
    records = []
    for elem in root:
        rec = dict(elem.attrib)
        for child in elem:
            rec[child.tag] = (child.text or "").strip()
        if not len(elem) and not elem.attrib:
            rec[elem.tag] = (elem.text or "").strip()
        records.append(rec)
    return _from_records(records)


def read_table(path: Path) -> Table:
    ext = normalize_ext(file_ext(path))
    if ext == "csv":
        return read_csv(path)
    if ext == "tsv":
        return read_csv(path, "\t")
    if ext == "json":
        return read_json(path)
    if ext == "xlsx":
        return read_xlsx(path)
    if ext == "xml":
        return read_xml(path)
    raise ValueError(ext)


# --- Запис --------------------------------------------------------------------

def _cell(v) -> str:
    return "" if v is None else str(v)


def _typed(v):
    """Рядок '42' -> 42, '3.14' -> 3.14 (для JSON та XLSX)."""
    if not isinstance(v, str):
        return v
    s = v.strip()
    if re.fullmatch(r"-?(0|[1-9]\d{0,14})", s):
        return int(s)
    if re.fullmatch(r"-?\d+\.\d+", s):
        try:
            return float(s)
        except ValueError:
            pass
    return v


def _xml_tag(name: str) -> str:
    tag = re.sub(r"[^\w.-]", "_", str(name).strip()) or "field"
    if not re.match(r"[A-Za-z_]", tag):
        tag = "_" + tag
    return tag


def write_table(table: Table, dst: Path, target: str) -> None:
    headers, rows = table
    if target in ("csv", "tsv"):
        # utf-8-sig, щоб Excel коректно відкрив кирилицю
        with open(dst, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f, delimiter="\t" if target == "tsv" else ",")
            w.writerow(headers)
            w.writerows([[_cell(c) for c in r] for r in rows])
    elif target == "json":
        records = [{h: _typed(r[i]) if i < len(r) else None for i, h in enumerate(headers)} for r in rows]
        dst.write_text(json.dumps(records, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    elif target == "xlsx":
        import openpyxl
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Дані"
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([_typed(c) for c in r])
        ws.freeze_panes = "A2"
        for i, h in enumerate(headers, 1):
            width = max([len(_cell(h))] + [len(_cell(r[i - 1])) for r in rows[:500] if i - 1 < len(r)])
            ws.column_dimensions[get_column_letter(i)].width = min(60, width + 2)
        wb.save(dst)
    elif target == "xml":
        root = ET.Element("rows")
        tags = [_xml_tag(h) for h in headers]
        for r in rows:
            row_el = ET.SubElement(root, "row")
            for i, tag in enumerate(tags):
                ET.SubElement(row_el, tag).text = _cell(r[i]) if i < len(r) else ""
        ET.indent(root)
        ET.ElementTree(root).write(dst, encoding="utf-8", xml_declaration=True)
    elif target == "html":
        esc = html_lib.escape
        head = "".join(f"<th>{esc(_cell(h))}</th>" for h in headers)
        body = "\n".join("<tr>" + "".join(f"<td>{esc(_cell(c))}</td>" for c in r) + "</tr>" for r in rows)
        dst.write_text(
            "<!DOCTYPE html>\n<html lang=\"uk\"><head><meta charset=\"utf-8\">"
            f"<title>{esc(dst.stem)}</title><style>"
            "body{font-family:Arial,sans-serif;margin:2em}table{border-collapse:collapse}"
            "th,td{border:1px solid #ccc;padding:4px 8px;text-align:left}th{background:#f0f0f0}"
            "tr:nth-child(even) td{background:#fafafa}</style></head><body>\n"
            f"<table>\n<thead><tr>{head}</tr></thead>\n<tbody>\n{body}\n</tbody>\n</table>\n</body></html>\n",
            encoding="utf-8")
    elif target == "md":
        def md(v):
            return _cell(v).replace("|", "\\|").replace("\n", " ")
        lines = ["| " + " | ".join(md(h) for h in headers) + " |",
                 "|" + "|".join("---" for _ in headers) + "|"]
        lines += ["| " + " | ".join(md(c) for c in r) + " |" for r in rows]
        dst.write_text("\n".join(lines) + "\n", encoding="utf-8")
    else:
        raise ValueError(target)


def convert(job, dst: Path, progress) -> None:
    table = read_table(job.src)
    progress(0.5)
    job.check_cancel()
    write_table(table, dst, job.target)
