"""Документи: PDF, DOCX, TXT, Markdown, HTML.

Усі формати читаються в спільну проміжну модель — список блоків
(заголовки, абзаци, пункти списку, код) з форматуванням тексту, —
а потім записуються в цільовий формат.
"""
from __future__ import annotations

import html as html_lib
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path

import pymupdf

from core import file_ext, normalize_ext


@dataclass
class Run:
    text: str
    bold: bool = False
    italic: bool = False


@dataclass
class Block:
    kind: str                       # h1..h6, p, li, pre
    runs: list[Run] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "".join(r.text for r in self.runs)


# --- Читання ------------------------------------------------------------------

def read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1251", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def blocks_from_txt(text: str) -> list[Block]:
    return [Block("p", [Run(line.rstrip())]) for line in text.splitlines() if line.strip()]


class _HtmlToBlocks(HTMLParser):
    BLOCK_TAGS = {"p", "div", "li", "pre", "h1", "h2", "h3", "h4", "h5", "h6",
                  "tr", "blockquote", "section", "article", "header", "footer", "dt", "dd"}
    SKIP_TAGS = {"script", "style", "head", "title", "noscript", "template"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocks: list[Block] = []
        self.current: Block | None = None
        self.kind_stack: list[str] = []
        self.bold = 0
        self.italic = 0
        self.skip = 0

    def _flush(self):
        if self.current and self.current.text.strip():
            if self.current.kind != "pre":
                # схлопуємо пробіли, як це робить браузер
                for r in self.current.runs:
                    r.text = re.sub(r"\s+", " ", r.text)
                self.current.runs[0].text = self.current.runs[0].text.lstrip()
                self.current.runs[-1].text = self.current.runs[-1].text.rstrip()
            self.blocks.append(self.current)
        self.current = None

    def _kind(self) -> str:
        for k in reversed(self.kind_stack):
            if k in ("h1", "h2", "h3", "h4", "h5", "h6", "li", "pre"):
                return k
        return "p"

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP_TAGS:
            self.skip += 1
        elif tag in self.BLOCK_TAGS:
            self._flush()
            self.kind_stack.append(tag)
        elif tag == "br":
            self._flush()
        elif tag in ("td", "th"):
            if self.current and self.current.text:
                self.current.runs.append(Run(" | "))
            if tag == "th":
                self.bold += 1
        elif tag in ("b", "strong"):
            self.bold += 1
        elif tag in ("i", "em"):
            self.italic += 1

    def handle_endtag(self, tag):
        if tag in self.SKIP_TAGS:
            self.skip = max(0, self.skip - 1)
        elif tag in self.BLOCK_TAGS:
            self._flush()
            if tag in self.kind_stack:
                idx = len(self.kind_stack) - 1 - self.kind_stack[::-1].index(tag)
                del self.kind_stack[idx:]
        elif tag in ("b", "strong", "th"):
            self.bold = max(0, self.bold - 1)  # th — заголовок клітинки, жирний
        elif tag in ("i", "em"):
            self.italic = max(0, self.italic - 1)

    def handle_data(self, data):
        if self.skip:
            return
        if self.current is None:
            if not data.strip():
                return
            self.current = Block(self._kind())
        self.current.runs.append(Run(data, bool(self.bold), bool(self.italic)))

    def close(self):
        super().close()
        self._flush()


def blocks_from_html(html: str) -> list[Block]:
    parser = _HtmlToBlocks()
    parser.feed(html)
    parser.close()
    return parser.blocks


def markdown_to_html(md_text: str) -> str:
    import markdown
    return markdown.markdown(md_text, extensions=["tables", "fenced_code", "sane_lists"])


def blocks_from_docx(path: Path) -> list[Block]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    blocks: list[Block] = []
    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            if not para.text.strip():
                continue
            style = (para.style.name or "").lower() if para.style is not None else ""
            m = re.match(r"heading (\d)", style)
            if m:
                kind = f"h{min(int(m.group(1)), 6)}"
            elif style == "title":
                kind = "h1"
            elif style == "subtitle":
                kind = "h2"
            elif "list" in style:
                kind = "li"
            else:
                kind = "p"
            runs = [Run(r.text, bool(r.bold), bool(r.italic)) for r in para.runs if r.text]
            blocks.append(Block(kind, runs or [Run(para.text)]))
        elif tag == "tbl":
            table = Table(child, document)
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                # об'єднані клітинки повторюються — прибираємо дублікати поспіль
                dedup = [c for i, c in enumerate(cells) if i == 0 or c != cells[i - 1]]
                line = " | ".join(dedup)
                if line.strip(" |"):
                    blocks.append(Block("p", [Run(line)]))
    return blocks


def blocks_from_pdf(path: Path, job=None) -> list[Block]:
    blocks: list[Block] = []
    with pymupdf.open(path) as doc:
        for page in doc:
            if job:
                job.check_cancel()
            for b in page.get_text("blocks", sort=True):
                if b[6] != 0:      # 0 = текстовий блок, 1 = зображення
                    continue
                text = re.sub(r"\s*\n\s*", " ", b[4]).strip()
                if text:
                    blocks.append(Block("p", [Run(text)]))
    return blocks


def read_blocks(path: Path, job=None) -> tuple[list[Block], str]:
    """Повертає (блоки, заголовок документа)."""
    ext = normalize_ext(file_ext(path))
    title = path.stem
    if ext == "txt":
        return blocks_from_txt(read_text_file(path)), title
    if ext == "md":
        return blocks_from_html(markdown_to_html(read_text_file(path))), title
    if ext == "html":
        raw = read_text_file(path)
        m = re.search(r"<title[^>]*>(.*?)</title>", raw, re.I | re.S)
        if m and m.group(1).strip():
            title = html_lib.unescape(m.group(1).strip())
        return blocks_from_html(raw), title
    if ext == "docx":
        return blocks_from_docx(path), title
    if ext == "pdf":
        return blocks_from_pdf(path, job), title
    raise ValueError(ext)


# --- Запис --------------------------------------------------------------------

def to_txt(blocks: list[Block]) -> str:
    out = []
    for b in blocks:
        text = b.text
        if b.kind == "li":
            text = "• " + text
        elif b.kind.startswith("h"):
            text = text.upper() if b.kind in ("h1", "h2") else text
        out.append(text)
    return "\n\n".join(out) + "\n"


def _md_runs(runs: list[Run]) -> str:
    parts = []
    for r in runs:
        t = r.text
        core = t.strip()
        if not core or not (r.bold or r.italic):
            parts.append(t)
            continue
        mark = "***" if r.bold and r.italic else "**" if r.bold else "*"
        lead = t[: len(t) - len(t.lstrip())]
        trail = t[len(t.rstrip()):]
        parts.append(f"{lead}{mark}{core}{mark}{trail}")
    return "".join(parts)


def to_markdown(blocks: list[Block]) -> str:
    out = []
    prev = None
    for b in blocks:
        if b.kind.startswith("h"):
            out.append("#" * int(b.kind[1]) + " " + b.text.strip())
        elif b.kind == "li":
            line = "- " + _md_runs(b.runs)
            if prev == "li":
                out[-1] += "\n" + line
                prev = b.kind
                continue
            out.append(line)
        elif b.kind == "pre":
            out.append("```\n" + b.text.rstrip("\n") + "\n```")
        else:
            out.append(_md_runs(b.runs))
        prev = b.kind
    return "\n\n".join(out) + "\n"


def _html_runs(runs: list[Run]) -> str:
    parts = []
    for r in runs:
        t = html_lib.escape(r.text)
        if r.italic:
            t = f"<em>{t}</em>"
        if r.bold:
            t = f"<strong>{t}</strong>"
        parts.append(t)
    return "".join(parts)


HTML_CSS = """
body { font-family: Arial, 'Segoe UI', sans-serif; font-size: 11pt; line-height: 1.45;
       color: #1f2328; max-width: 820px; margin: 2em auto; padding: 0 1em; }
h1, h2, h3, h4, h5, h6 { line-height: 1.25; margin: 1.1em 0 0.5em; }
p, li { margin: 0.35em 0; }
pre { font-family: Consolas, 'Courier New', monospace; background: #f4f4f4;
      padding: 0.7em; white-space: pre-wrap; }
"""


def to_html(blocks: list[Block], title: str) -> str:
    body = []
    in_list = False
    for b in blocks:
        if b.kind == "li" and not in_list:
            body.append("<ul>")
            in_list = True
        elif b.kind != "li" and in_list:
            body.append("</ul>")
            in_list = False
        if b.kind == "pre":
            body.append(f"<pre>{html_lib.escape(b.text)}</pre>")
        else:
            body.append(f"<{b.kind}>{_html_runs(b.runs)}</{b.kind}>")
    if in_list:
        body.append("</ul>")
    return (
        "<!DOCTYPE html>\n<html lang=\"uk\">\n<head>\n<meta charset=\"utf-8\">\n"
        f"<title>{html_lib.escape(title)}</title>\n<style>{HTML_CSS}</style>\n</head>\n<body>\n"
        + "\n".join(body) + "\n</body>\n</html>\n"
    )


def write_docx(blocks: list[Block], dst: Path) -> None:
    import docx
    from docx.shared import Pt

    document = docx.Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for b in blocks:
        if b.kind.startswith("h"):
            para = document.add_heading(level=min(int(b.kind[1]), 9))
        elif b.kind == "li":
            para = document.add_paragraph(style="List Bullet")
        else:
            para = document.add_paragraph()
        for r in b.runs:
            run = para.add_run(r.text)
            run.bold = r.bold or None
            run.italic = r.italic or None
            if b.kind == "pre":
                run.font.name = "Consolas"
    document.save(str(dst))


def _font_archive() -> tuple[pymupdf.Archive, str]:
    """Шрифт з кирилицею для PDF: системний Arial/Segoe UI, інакше вбудований."""
    fonts_dir = Path(r"C:\Windows\Fonts")
    for regular, bold, italic, bolditalic in (
        ("arial.ttf", "arialbd.ttf", "ariali.ttf", "arialbi.ttf"),
        ("segoeui.ttf", "segoeuib.ttf", "segoeuii.ttf", "segoeuiz.ttf"),
        ("DejaVuSans.ttf", "DejaVuSans-Bold.ttf", "DejaVuSans-Oblique.ttf", "DejaVuSans-BoldOblique.ttf"),
    ):
        if (fonts_dir / regular).exists():
            css = ""
            for fname, weight, style in ((regular, "normal", "normal"), (bold, "bold", "normal"),
                                         (italic, "normal", "italic"), (bolditalic, "bold", "italic")):
                if (fonts_dir / fname).exists():
                    css += (f"@font-face {{ font-family: DocFont; src: url({fname}); "
                            f"font-weight: {weight}; font-style: {style}; }}\n")
            css += "body, p, li, h1, h2, h3, h4, h5, h6, td, th, div, span { font-family: DocFont; }\n"
            return pymupdf.Archive(str(fonts_dir)), css
    return pymupdf.Archive(), ""


def html_to_pdf(html: str, dst: Path, job=None) -> None:
    archive, font_css = _font_archive()
    css = font_css + """
        body { font-size: 11pt; line-height: 1.4; }
        h1 { font-size: 20pt; } h2 { font-size: 16pt; } h3 { font-size: 13pt; }
        pre { font-family: monospace; font-size: 9pt; }
    """
    story = pymupdf.Story(html=html, user_css=css, archive=archive)
    page_rect = pymupdf.paper_rect("a4")
    content_rect = page_rect + (56, 56, -56, -56)  # поля ~2 см
    writer = pymupdf.DocumentWriter(str(dst))
    more = True
    while more:
        if job:
            job.check_cancel()
        device = writer.begin_page(page_rect)
        more, _ = story.place(content_rect)
        story.draw(device)
        writer.end_page()
    writer.close()

    # вбудовуємо лише використані гліфи шрифту — інакше PDF важить мегабайти
    with pymupdf.open(dst) as doc:
        doc.subset_fonts()
        data = doc.tobytes(garbage=3, deflate=True)
    dst.write_bytes(data)


def pdf_to_images(job, dst: Path, progress) -> Path:
    fmt = job.target
    zoom = 2.0  # ~144 DPI
    with pymupdf.open(job.src) as doc:
        n = doc.page_count
        if n == 1:
            targets = [dst]
        else:
            folder = dst.with_name(dst.name[: -len(fmt) - 1] + f" ({fmt})")
            i = 1
            while folder.exists():
                folder = dst.with_name(dst.name[: -len(fmt) - 1] + f" ({fmt} {i})")
                i += 1
            folder.mkdir(parents=True)
            targets = [folder / f"page_{p + 1:03d}.{fmt}" for p in range(n)]
        for p, page in enumerate(doc):
            job.check_cancel()
            pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False)
            if fmt == "jpg":
                pix.save(str(targets[p]), jpg_quality=job.options.image_quality)
            else:
                pix.save(str(targets[p]))
            progress((p + 1) / n)
    return targets[0] if n == 1 else targets[0].parent


def convert(job, dst: Path, progress) -> Path | None:
    src_ext = normalize_ext(file_ext(job.src))
    target = job.target

    if src_ext == "pdf" and target in ("png", "jpg"):
        return pdf_to_images(job, dst, progress)

    blocks, title = read_blocks(job.src, job)
    progress(0.5)
    job.check_cancel()

    if target == "txt":
        dst.write_text(to_txt(blocks), encoding="utf-8")
    elif target == "md":
        dst.write_text(to_markdown(blocks), encoding="utf-8")
    elif target == "html":
        dst.write_text(to_html(blocks, title), encoding="utf-8")
    elif target == "docx":
        write_docx(blocks, dst)
    elif target == "pdf":
        if src_ext == "html":
            source_html = read_text_file(job.src)    # зберігаємо оригінальну верстку
        elif src_ext == "md":
            source_html = markdown_to_html(read_text_file(job.src))
        else:
            source_html = to_html(blocks, title)
        html_to_pdf(source_html, dst, job)
    else:
        raise ValueError(target)
    return None
